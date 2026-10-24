"""
Builds dashboard/ecommerce_dashboard.xlsx - an Excel MIS dashboard.

Every number on the Dashboard sheet is a live formula (SUMIFS / COUNTIFS / INDEX-MATCH) over the
data sheets, so the dashboard updates if the data changes. Only the RFM scores are pasted values
(they come from the SQL view vw_customer_rfm, recomputed in scripts/analysis.py).

Usage:  python scripts/make_dashboard.py
"""
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import analysis

OUT = analysis.ROOT / "dashboard" / "ecommerce_dashboard.xlsx"

FONT = "Arial"
DARK = "184F95"
BLUE = "2A78D6"
ORANGE = "EB6834"
LIGHT = "EAF2FC"
GREY = "52514E"
RUPEE = '"₹"#,##0'
PCT = "0.0%"

thin = Side(style="thin", color="D9D9D9")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def font(**kw):
    return Font(name=FONT, **kw)


def header_row(ws, row, col, labels):
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=label)
        c.font = font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=DARK)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER


def body_cell(ws, row, col, value, fmt=None, bold=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = font(bold=bold)
    c.border = BORDER
    if fmt:
        c.number_format = fmt
    return c


def section_title(ws, row, text):
    c = ws.cell(row=row, column=2, value=text)
    c.font = font(bold=True, size=13, color=DARK)


def data_sheet(wb, title, df, formulas=None, formats=None):
    """Write a dataframe as a data sheet, then add formula columns {header: template with {r}}."""
    ws = wb.create_sheet(title)
    cols = list(df.columns) + list((formulas or {}).keys())
    header_row(ws, 1, 1, cols)
    for r, row in enumerate(df.itertuples(index=False), start=2):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v).font = font()
        for c, tmpl in enumerate((formulas or {}).values(), start=len(df.columns) + 1):
            ws.cell(row=r, column=c, value=tmpl.format(r=r)).font = font()
    for header, fmt in (formats or {}).items():
        col = cols.index(header) + 1
        for r in range(2, len(df) + 2):
            ws.cell(row=r, column=col).number_format = fmt
    for i, header in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(header) + 4)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(df) + 1}"
    return ws


def style_chart(chart, color, title, horizontal=False):
    chart.title = title
    chart.title.overlay = False
    chart.legend = None
    chart.height = 7.2
    chart.width = 16
    # openpyxl hides both axes unless told otherwise
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    s = chart.series[0]
    s.graphicalProperties.solidFill = color
    s.graphicalProperties.line.solidFill = color
    if horizontal:
        chart.type = "bar"
        chart.y_axis.majorGridlines = None
        chart.x_axis.scaling.orientation = "maxMin"  # first row on top
        chart.y_axis.delete = True
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True
        chart.dataLabels.showSerName = False
        chart.dataLabels.showCatName = False
        chart.dataLabels.showLegendKey = False
        chart.gapWidth = 60
    return chart


def main():
    customers, products, orders, details = analysis.load()
    n_c, n_p, n_o, n_d = len(customers) + 1, len(products) + 1, len(orders) + 1, len(details) + 1

    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"

    # ---------------- data sheets with helper formula columns ----------------
    data_sheet(wb, "Customers", customers, {
        "total_orders": f"=COUNTIF(Orders!$C$2:$C${n_o},A{{r}})",
        "total_spent": f"=SUMIF(Orders!$C$2:$C${n_o},A{{r}},Orders!$D$2:$D${n_o})",
        "segment": '=IF(D{r}=0,"Never ordered",IF(D{r}=1,"One-time buyer",IF(D{r}<=4,"Occasional shopper","Regular customer")))',
    }, {"total_spent": RUPEE})
    data_sheet(wb, "Products", products, formats={"price": RUPEE})
    o = orders.copy()
    o["order_date"] = o["order_date"].dt.date
    ws_o = data_sheet(wb, "Orders", o, {
        "location": f"=INDEX(Customers!$C$2:$C${n_c},MATCH(C{{r}},Customers!$A$2:$A${n_c},0))",
        "month_key": "=YEAR(B{r})*100+MONTH(B{r})",
    }, {"total_amount": RUPEE})
    for r in range(2, n_o + 1):
        ws_o.cell(row=r, column=2).number_format = "yyyy-mm-dd"
    data_sheet(wb, "OrderDetails", details, {
        "line_amount": "=C{r}*D{r}",
        "product_name": f"=INDEX(Products!$B$2:$B${n_p},MATCH(B{{r}},Products!$A$2:$A${n_p},0))",
        "category": f"=INDEX(Products!$C$2:$C${n_p},MATCH(B{{r}},Products!$A$2:$A${n_p},0))",
        "month_key": f"=INDEX(Orders!$F$2:$F${n_o},MATCH(A{{r}},Orders!$A$2:$A${n_o},0))",
        "location": f"=INDEX(Orders!$E$2:$E${n_o},MATCH(A{{r}},Orders!$A$2:$A${n_o},0))",
    }, {"price_per_unit": RUPEE, "line_amount": RUPEE})

    # RFM: pasted values from the SQL view + live summary
    rfm = analysis.rfm(customers, orders)[[
        "customer_id", "name", "location", "recency_days", "frequency", "monetary",
        "r_score", "f_score", "m_score", "rfm_segment"]].sort_values("monetary", ascending=False)
    ws_r = data_sheet(wb, "RFM", rfm, formats={"monetary": RUPEE})
    n_r = len(rfm) + 1
    ws_r["L1"] = "Source: values from SQL view vw_customer_rfm (sql/03_advanced_sql.sql), scores 1-4 via NTILE(4)."
    ws_r["L1"].font = font(italic=True, color=GREY)
    header_row(ws_r, 3, 12, ["Segment", "Customers", "Revenue", "Revenue share"])
    segments = ["Champions", "Loyal / Potential", "Needs Attention", "At Risk", "Lost"]
    for i, seg in enumerate(segments, start=4):
        body_cell(ws_r, i, 12, seg)
        body_cell(ws_r, i, 13, f"=COUNTIF($J$2:$J${n_r},L{i})")
        body_cell(ws_r, i, 14, f"=SUMIF($J$2:$J${n_r},L{i},$F$2:$F${n_r})", RUPEE)
        body_cell(ws_r, i, 15, f"=N{i}/SUM($N$4:$N$8)", PCT)
    for col, w in zip("LMNO", (20, 12, 16, 14)):
        ws_r.column_dimensions[col].width = w

    # ---------------- dashboard ----------------
    dash.sheet_view.showGridLines = False
    dash.column_dimensions["A"].width = 2
    for col in "BCDEFG":
        dash.column_dimensions[col].width = 15
    dash.column_dimensions["B"].width = 26

    dash["B1"] = "E-Commerce Sales Dashboard"
    dash["B1"].font = font(bold=True, size=20, color=DARK)
    dash["B2"] = ("Orders from Mar 2023 to Feb 2024  ·  all figures are live formulas over the "
                  "Customers, Products, Orders and OrderDetails sheets")
    dash["B2"].font = font(italic=True, color=GREY)

    # KPI tiles: label row 4, value row 5
    kpis = [
        ("Total Revenue", f"=SUM(Orders!$D$2:$D${n_o})", RUPEE),
        ("Total Orders", f"=COUNT(Orders!$A$2:$A${n_o})", "#,##0"),
        ("Avg Order Value", "=B5/C5", RUPEE),
        ("Total Customers", f"=COUNTA(Customers!$A$2:$A${n_c})", "#,##0"),
        ("Buying Customers", f'=COUNTIF(Customers!$D$2:$D${n_c},">0")', "#,##0"),
        ("Repeat Rate", f'=COUNTIF(Customers!$D$2:$D${n_c},">1")/F5', PCT),
    ]
    for i, (label, formula, fmt) in enumerate(kpis):
        col = 2 + i
        top = dash.cell(row=4, column=col, value=label)
        top.font = font(bold=True, color="FFFFFF", size=9)
        top.fill = PatternFill("solid", fgColor=BLUE)
        top.alignment = Alignment(horizontal="center")
        val = dash.cell(row=5, column=col, value=formula)
        val.font = font(bold=True, size=16, color=DARK)
        val.fill = PatternFill("solid", fgColor=LIGHT)
        val.alignment = Alignment(horizontal="center", vertical="center")
        val.number_format = fmt
    dash.row_dimensions[5].height = 34

    # Monthly sales
    months = sorted(orders["order_date"].dt.strftime("%Y-%m").unique())
    start = 8
    section_title(dash, start, "Monthly sales trend")
    header_row(dash, start + 1, 2, ["Month", "Month key", "Orders", "Sales", "MoM change"])
    for i, m in enumerate(months):
        r = start + 2 + i
        body_cell(dash, r, 2, m)
        body_cell(dash, r, 3, int(m.replace("-", ""))).font = font(color=GREY)
        body_cell(dash, r, 4, f"=COUNTIF(Orders!$F$2:$F${n_o},C{r})")
        body_cell(dash, r, 5, f"=SUMIF(Orders!$F$2:$F${n_o},C{r},Orders!$D$2:$D${n_o})", RUPEE)
        body_cell(dash, r, 6, "" if i == 0 else f"=(E{r}-E{r - 1})/E{r - 1}", PCT)
    last = start + 1 + len(months)
    body_cell(dash, last + 1, 2, "Total", bold=True)
    body_cell(dash, last + 1, 4, f"=SUM(D{start + 2}:D{last})", bold=True)
    body_cell(dash, last + 1, 5, f"=SUM(E{start + 2}:E{last})", RUPEE, bold=True)
    chart = LineChart()
    chart.add_data(Reference(dash, min_col=5, min_row=start + 1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(dash, min_col=2, min_row=start + 2, max_row=last))
    style_chart(chart, BLUE, "Monthly sales (₹)")
    chart.series[0].smooth = False
    chart.y_axis.numFmt = '#,##0'
    dash.add_chart(chart, f"H{start}")

    def breakdown(start, title, labels, cols, chart_title, value_col, color=BLUE):
        """Small table (label + formula columns) with a horizontal bar chart beside it."""
        section_title(dash, start, title)
        header_row(dash, start + 1, 2, [c[0] for c in cols])
        for i, label in enumerate(labels):
            r = start + 2 + i
            body_cell(dash, r, 2, label)
            for j, (_, tmpl, fmt) in enumerate(cols[1:], start=3):
                body_cell(dash, r, j, tmpl.format(r=r), fmt)
        end = start + 1 + len(labels)
        ch = BarChart()
        ch.add_data(Reference(dash, min_col=value_col, min_row=start + 1, max_row=end), titles_from_data=True)
        ch.set_categories(Reference(dash, min_col=2, min_row=start + 2, max_row=end))
        style_chart(ch, color, chart_title, horizontal=True)
        ch.dataLabels.numFmt = '#,##0'
        dash.add_chart(ch, f"H{start}")

    od = f"OrderDetails!$A$2:$A${n_d}"
    od_col = lambda c: od.replace("$A$", f"${c}$")

    # Category (sorted by revenue so the chart reads top-down)
    lines = analysis.order_lines(customers, products, orders, details)
    cats = lines.groupby("category")["line_amount"].sum().sort_values(ascending=False).index.tolist()
    breakdown(25, "Revenue by category", cats, [
        ("Category", None, None),
        ("Units sold", f"=SUMIF({od_col('G')},B{{r}},{od_col('C')})", "#,##0"),
        ("Revenue", f"=SUMIF({od_col('G')},B{{r}},{od_col('E')})", RUPEE),
        ("Share", "=D{r}/SUM($D$27:$D$29)", PCT),
    ], "Revenue by category (₹)", 4)

    prods = lines.groupby("product_name")["line_amount"].sum().sort_values(ascending=False).index.tolist()
    breakdown(42, "Revenue by product", prods, [
        ("Product", None, None),
        ("Units sold", f"=SUMIF({od_col('F')},B{{r}},{od_col('C')})", "#,##0"),
        ("Revenue", f"=SUMIF({od_col('F')},B{{r}},{od_col('E')})", RUPEE),
        ("Order lines", f"=COUNTIF({od_col('F')},B{{r}})", "#,##0"),
    ], "Revenue by product (₹)", 4)

    cities = lines.groupby("location")["line_amount"].sum().sort_values(ascending=False).index.tolist()
    breakdown(59, "Revenue by city", cities, [
        ("City", None, None),
        ("Customers", f"=COUNTIF(Customers!$C$2:$C${n_c},B{{r}})", "#,##0"),
        ("Orders", f"=COUNTIF(Orders!$E$2:$E${n_o},B{{r}})", "#,##0"),
        ("Revenue", f"=SUMIF(Orders!$E$2:$E${n_o},B{{r}},Orders!$D$2:$D${n_o})", RUPEE),
    ], "Revenue by city (₹)", 5)

    breakdown(76, "Customer segments by number of orders",
              ["Never ordered", "One-time buyer", "Occasional shopper", "Regular customer"], [
        ("Segment", None, None),
        ("Customers", f"=COUNTIF(Customers!$F$2:$F${n_c},B{{r}})", "#,##0"),
        ("Revenue", f"=SUMIF(Customers!$F$2:$F${n_c},B{{r}},Customers!$E$2:$E${n_c})", RUPEE),
        ("Rule", '=IF(B{r}="Never ordered","0 orders",IF(B{r}="One-time buyer","1 order",'
                 'IF(B{r}="Occasional shopper","2-4 orders","5+ orders")))', None),
    ], "Customers per segment", 3, ORANGE)

    breakdown(93, "RFM segments (from SQL view vw_customer_rfm)", segments, [
        ("RFM segment", None, None),
        ("Customers", "=INDEX(RFM!$M$4:$M$8,MATCH(B{r},RFM!$L$4:$L$8,0))", "#,##0"),
        ("Revenue", "=INDEX(RFM!$N$4:$N$8,MATCH(B{r},RFM!$L$4:$L$8,0))", RUPEE),
        ("Revenue share", "=INDEX(RFM!$O$4:$O$8,MATCH(B{r},RFM!$L$4:$L$8,0))", PCT),
    ], "Revenue by RFM segment (₹)", 4)

    wb.calculation.fullCalcOnLoad = True
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print("saved", OUT.relative_to(analysis.ROOT))


if __name__ == "__main__":
    main()
